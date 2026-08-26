"""Kanban MES — web app (FastAPI + Jinja2 + HTMX).

Fluxo: /capture (foto ou folha manual) → staging local → /sheet/{uid} revisão
com células coloridas → /validate = única porta para o Postgres.
"""

from __future__ import annotations

import csv
import datetime
import hashlib
import io
import re
import threading
import time
from contextlib import asynccontextmanager
from dataclasses import asdict
from pathlib import Path

from fastapi import FastAPI, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from .. import db, imaging, pg_store
from ..config import settings
from ..matching import carryover, header_cross, loaders, operador
from ..matching.cross_check import check_sheet
from ..matching.params import CrossParams
from ..matching.scorer import Scorer
from ..ocr.provider import OcrError, empty_extraction, get_provider, rescue_header
from ..templates_spec import TEMPLATES, field_value, get_template, is_marked
from . import estado as estado_data
from . import export as cpis_export
from . import pdf as pdf_gen

class NoCacheStaticFiles(StaticFiles):
    """Estáticos com revalidação obrigatória — o link leva ?v=<hash>, e isto
    impede o browser/edge da Cloudflare de servir CSS velho a quem tem o URL antigo."""

    async def get_response(self, path: str, scope):
        response = await super().get_response(path, scope)
        response.headers["Cache-Control"] = "no-cache"
        return response


_STATIC_DIR = Path(__file__).parent / "static"


@asynccontextmanager
async def _lifespan(app: FastAPI):
    """No arranque, re-enfileirar folhas que um restart deixou em `pending`:
    as threads do worker são daemon e morrem com o processo — sem isto, um
    deploy a meio de um lote deixava folhas em spinner eterno."""
    try:
        conn = db.connect()
        try:
            stuck = db.pending_with_image(conn)
        finally:
            conn.close()
        if stuck:
            print(f"[arranque] a retomar OCR de {len(stuck)} folha(s) pendente(s)", flush=True)
            threading.Thread(target=_process_batch, args=(stuck,), daemon=True).start()
    except Exception as exc:
        print(f"[arranque] retoma de pendentes falhou: {exc}", flush=True)
    yield


app = FastAPI(title="Kanban MES", lifespan=_lifespan)
app.mount("/static", NoCacheStaticFiles(directory=str(_STATIC_DIR)), name="static")
templates = Jinja2Templates(directory=str(Path(__file__).parent / "templates"))
templates.env.globals["css_version"] = hashlib.sha1(
    (_STATIC_DIR / "design.css").read_bytes()
).hexdigest()[:10]
# a folha decide o que é uma marca; o template não repete a regra
templates.env.globals["is_marked"] = is_marked
# OF/OV apresentam-se como números puros (convenção do planeamento)
from ..matching import similarity as _sim  # noqa: E402
templates.env.globals["strip_ref"] = _sim.strip_ref_prefix
# lê a célula pelo nome atual e pelo antigo (folhas lidas antes do rename)
templates.env.globals["field_value"] = field_value


@app.middleware("http")
async def _attach_watermark(request: Request, call_next):
    """Marca de água = folha mais recente no momento do pedido.

    A página do histórico guarda o valor do primeiro paint e compara-o com o
    header de cada poll HTMX; a diferença são as folhas que entraram entretanto,
    e é isso que alimenta o banner «N folhas novas».
    """
    # Estáticos e fotos não precisam de marca de água — e abrir uma ligação
    # SQLite (com executescript do schema) por cada imagem servida punha o
    # event loop a pagar até 5 s de busy_timeout em cada pedido.
    path = request.url.path
    if path.startswith("/static/") or path.endswith("/photo"):
        return await call_next(request)
    watermark = 0
    if request.method == "GET":
        try:
            conn = _conn()
            try:
                row = conn.execute("SELECT MAX(rowid) AS m FROM sheets").fetchone()
                watermark = int(row["m"] or 0)
            finally:
                conn.close()
        except Exception:
            watermark = 0
    request.state.watermark = watermark
    response = await call_next(request)
    response.headers["X-Sheet-Watermark"] = str(watermark)
    return response


def tunnel_url() -> str | None:
    """URL público atual do túnel Cloudflare (escrito por scripts/tunnel.sh)."""
    try:
        text = (settings.data_dir / "tunnel_url.txt").read_text().strip()
        return text.splitlines()[-1] if text else None
    except OSError:
        return None

# De quanto em quanto tempo se pergunta ao Postgres se há plano novo. Baixo
# porque a pergunta é uma linha; o índice só se reconstrói se a resposta mudar.
_FRESHNESS_PROBE_SECONDS = 30
_index_cache: dict[str, tuple[float, object, str | None]] = {}
_index_lock = threading.Lock()


def _conn():
    return db.connect()


def get_index(loader_name: str):
    """Índice do plano em cache, revalidado contra o snapshot mais recente.

    O TTL sozinho era a forma errada de o fazer: reconstruía 64 mil linhas de
    dez em dez minutos mesmo sem nada ter mudado, e ainda assim demorava até
    dez minutos a ver um plano novo. A sonda é uma linha de SQL — reconstrói
    quando (e só quando) o snapshot muda.
    """
    now = time.monotonic()
    with _index_lock:
        hit = _index_cache.get(loader_name)
    if hit:
        checked_at, index, snapshot = hit
        if now - checked_at < _FRESHNESS_PROBE_SECONDS:
            return index
        current = _current_index_snapshot(loader_name)
        if current is None or current == snapshot:
            # sonda falhou (Postgres em baixo) ou nada mudou: continuar com o
            # que temos, e voltar a sondar daqui a pouco
            with _index_lock:
                _index_cache[loader_name] = (time.monotonic(), index, snapshot)
            return index
    index = getattr(loaders, loader_name)()
    with _index_lock:
        _index_cache[loader_name] = (
            time.monotonic(), index, _current_index_snapshot(loader_name)
        )
    return index


def _current_snapshot_id() -> str | None:
    try:
        return (loaders.plan_snapshot_info() or {}).get("snapshot_id")
    except Exception:
        return None


def _current_index_snapshot(loader_name: str) -> str | None:
    """Cada índice invalida-se pela SUA carga: os colaboradores chegam num
    snapshot próprio e ficavam presos ao snapshot do plano."""
    if loader_name == "load_employees":
        try:
            return loaders.employees_snapshot_id()
        except Exception:
            return None
    return _current_snapshot_id()


def get_employees():
    """Colaboradores em cache, invalidada pela sua própria carga."""
    return get_index("load_employees")


def make_scorer(template_name: str) -> Scorer:
    template = get_template(template_name)
    index = get_index(template.index_loader)
    active = loaders.load_active_ofs() if template.family == "cantoneiras" else set()
    return Scorer(index, CrossParams.load(), active_primary=active)


_header_machine_cache: tuple[float, list] | None = None


def _load_header_machines() -> list:
    """Catálogo pequeno de máquinas, isolado do índice pesado do plano.

    Uma fonte em baixo devolve vazio: o checker marca ``no_reference`` e nunca
    inventa.
    """
    global _header_machine_cache
    now = time.monotonic()
    if _header_machine_cache and now - _header_machine_cache[0] < 60:
        return _header_machine_cache[1]
    try:
        values = list(loaders.load_machines())
    except Exception:
        values = []
    _header_machine_cache = (now, values)
    return values


def _source_document(sheet: dict) -> dict:
    """Proveniência do ficheiro, nunca uma referência para a data da folha."""
    filename, page = pg_store.source_from_image_path(sheet.get("image_path"))
    if not filename:
        return {}
    return {
        "filename": filename,
        "page": page,
        # derivado do nome do render, não guardado à parte — por isso marcado
        "inferred": True,
        "date_is_provenance_only": True,
    }


# Prefixo dd-mm-aaaa que o scanner da fábrica põe no nome dos PDFs
# («06-08-2026 - Rapid20T 2.pdf»).
_SOURCE_DATE_RE = re.compile(r"^(\d{2})-(\d{2})-(\d{4})")


def _assumed_sheet_date(sheet: dict) -> str | None:
    """Data assumida da folha (dd/mm/aaaa): dia útil anterior à data-base.

    Regra da fábrica (26/08): as folhas entregues à digitalização são sempre
    do dia útil anterior. A data-base é o prefixo dd-mm-aaaa do nome do PDF de
    origem; sem PDF (foto, folha manual), vale o `created_at` da folha.
    """
    filename, _page = pg_store.source_from_image_path(sheet.get("image_path"))
    base: datetime.date | None = None
    if filename:
        m = _SOURCE_DATE_RE.match(filename)
        if m:
            try:
                base = datetime.date(int(m.group(3)), int(m.group(2)),
                                     int(m.group(1)))
            except ValueError:
                base = None
    if base is None:
        try:
            base = datetime.datetime.fromisoformat(
                str(sheet.get("created_at") or "")).date()
        except ValueError:
            return None
    return header_cross.previous_business_day(base).strftime("%d/%m/%Y")


def _plan_header_machines(cross: dict, scorer: Scorer | None) -> list[str]:
    """Máquinas não vazias das linhas que ligaram fortemente ao plano."""
    if scorer is None:
        return []
    key_name = scorer.index.spec.key_field
    by_key = {
        str(entry.get(key_name)): entry
        for entry in scorer.index.entries
        if entry.get(key_name) is not None
    }
    out: list[str] = []
    for row in cross.get("rows", []):
        if row.get("mode") != "strong" or not row.get("matched_plan_key"):
            continue
        entry = by_key.get(str(row["matched_plan_key"])) or {}
        value = str(entry.get("maquina") or "").strip()
        if value and value not in out:
            out.append(value)
    return out


def resolve_operator(conn, uid: str, sheet: dict) -> dict | None:
    """Resolve o operador da folha contra a lista de colaboradores.

    Corre para TODAS as folhas, incluindo o verso (paragens): é lá que estão
    metade dos casos, e a mesma pessoa aparecia com nomes diferentes na frente
    e no verso da mesma folha física.
    """
    header = (sheet["sheet_data"] or {}).get("header") or {}
    try:
        employees = get_employees()
    except Exception:
        return None
    if not employees:
        return None

    match = operador.resolve(header.get("operador"), header.get("n_operador"), employees)
    if match.cod is None and match.pernr is None:
        return asdict(match)

    protegidos = db.human_header_fields(conn, uid)
    updates = [
        (field, value)
        for field, value in (("operador", match.name), ("n_operador", str(match.cod or "")))
        if value and field not in protegidos
        and str(header.get(field) or "").strip() != value
    ]
    if updates:
        # Aplicar sobre o estado FRESCO, nunca sobre a cópia com que se
        # calculou: substituir o header inteiro por uma cópia velha apagava
        # edições humanas feitas enquanto o worker corria. Se a folha mudou
        # entretanto, desiste-se — a edição que a mudou dispara novo ciclo.
        fresh = db.get_sheet(conn, uid)
        if fresh and fresh["sheet_data"] and fresh["revision"] == sheet["revision"]:
            data = fresh["sheet_data"]
            if db.save_sheet_data(
                conn, uid,
                {**data, "header": {**(data.get("header") or {}),
                                    **dict(updates)}},
                fresh["revision"],
            ):
                for field, value in updates:
                    db.record_edit(conn, uid, f"header.{field}",
                                   (data.get("header") or {}).get(field), value,
                                   "system", "colaboradores")
    return asdict(match)


def run_cross_check(conn, uid: str) -> None:
    sheet = db.get_sheet(conn, uid)
    if not sheet or not sheet["sheet_data"]:
        return
    # O operador resolve-se primeiro e escreve como sempre escreveu (nome/nº
    # canónicos do SAP). O checker do cabeçalho corre DEPOIS, sobre o estado
    # já canónico: as células dele descrevem e propõem, nunca disputam a
    # decisão do resolve_operator.
    operator_match = resolve_operator(conn, uid, sheet)
    # Reler DEPOIS do resolve_operator (que pode ter escrito): o cruzamento é
    # calculado sobre uma revisão conhecida e só se grava se a folha ainda for
    # essa — um cross calculado sobre valores velhos não pode pintar os novos.
    base = db.get_sheet(conn, uid)
    if not base or not base["sheet_data"]:
        return
    template = get_template(base["template_name"])
    data = base["sheet_data"]
    rows = data.get("rows") or []
    header = data.get("header") or {}
    human_header = db.human_header_fields(conn, uid)
    expected = base["revision"]

    scorer: Scorer | None = None
    if template.index_loader is None:
        # ex.: paragens — não há plano contra que cruzar, mas o cabeçalho
        # (operador, máquina, data, turno) verifica-se na mesma.
        plan_reference = {"status": "not_applicable"}
    else:
        try:
            scorer = make_scorer(base["template_name"])
            plan_reference = {"status": "available"}
        except Exception:
            # O plano é uma fonte independente. Uma indisponibilidade não pode
            # impedir data/turno, colaboradores ou máquina de serem cruzados
            # e persistidos.
            plan_reference = {
                "status": "no_reference",
                "message": "Plano indisponível; linhas mantidas sem cruzamento.",
            }
    human_rows = db.human_fields_by_row(conn, uid)
    if scorer is not None:
        cross = check_sheet(rows, scorer, human_rows, footer=data.get("footer"))
    else:
        cross = {"summary": {}, "review_order": [], "rows": []}

    # Cabeçalho determinístico: cada fonte só entra se puder ser precisa —
    # falhas viram ``no_reference``, nunca palpites.
    employees = None
    if any(str(header.get(f) or "").strip() for f in ("operador", "n_operador")):
        try:
            employees = get_employees()
        except Exception:
            employees = None
    plan_machines = _plan_header_machines(cross, scorer)
    machine_catalog: list = []
    if (str(header.get("setor_maquina") or "").strip() or plan_machines
            or header_cross.template_machine(template)):
        machine_catalog = _load_header_machines()
    source_document = _source_document(base)
    # A data assumida também vale para o verso (paragens): é a mesma folha
    # física, digitalizada no mesmo dia.
    assumed_date = _assumed_sheet_date(base)

    def check_current_header() -> dict:
        return header_cross.check_header(
            header, template,
            human_fields=human_header,
            employees=employees,
            machines=machine_catalog,
            plan_machines=plan_machines,
            source_document=source_document,
            assumed_date=assumed_date,
        )

    header_result = check_current_header()

    # escrita automática (política de perda esperada), auditada como 'system';
    # valores, auditoria e cross final vão num único commit CAS
    edits: list[tuple[str, object, object, str]] = []
    applied_header: dict[str, dict] = {}
    for field_name, cell in header_result["cells"].items():
        proposal = cell.get("proposal")
        if not cell.get("auto_write") or proposal is None:
            continue
        old = header.get(field_name)
        if str(old or "").strip() == str(proposal).strip():
            continue
        edits.append((f"header.{field_name}", old, str(proposal),
                      cell.get("actor") or "header:cross"))
        header[field_name] = str(proposal)
        applied_header[field_name] = cell
    data["header"] = header
    if applied_header:
        # Recalcular sobre o estado final: as células têm de descrever o valor
        # que fica gravado, não o que existia antes da substituição.
        header_result = check_current_header()
        for field_name, original in applied_header.items():
            final_cell = header_result["cells"].get(field_name)
            if final_cell is None:
                continue
            final_cell["applied"] = True
            final_cell["actor"] = original.get("actor")
            final_cell["message"] = (
                "Substituído automaticamente. " + final_cell["message"]
            )

    applied_cells: list[tuple[int, str]] = []
    for rc in cross["rows"]:
        for cell in rc["cells"]:
            if cell["auto_write"] and cell["proposal"] is not None:
                i, f = rc["row_index"], cell["field"]
                if i < len(rows) and str(rows[i].get(f) or "").strip() != cell["proposal"]:
                    edits.append((f"rows[{i}].{f}", rows[i].get(f),
                                  cell["proposal"], "cross"))
                    rows[i][f] = cell["proposal"]
                    applied_cells.append((i, f))
    if applied_cells and scorer is not None:
        # Como no cabeçalho: recalcular sobre o estado final, para cada célula
        # descrever o valor que ficou gravado (não o que existia antes da
        # substituição), com a marca `applied` a dizer que foi o motor.
        cross_final = check_sheet(rows, scorer, human_rows,
                                  footer=data.get("footer"))
        final_by_row = {r["row_index"]: r for r in cross_final["rows"]}
        for i, f in applied_cells:
            final_cell = next(
                (c for c in final_by_row.get(i, {}).get("cells", [])
                 if c["field"] == f), None)
            if final_cell is not None:
                final_cell["applied"] = True
                final_cell["message"] = ("Substituído automaticamente. "
                                         "O valor lido pelo OCR fica visível na célula.")
        cross["rows"] = cross_final["rows"]
        cross["summary"] = cross_final["summary"]
        cross["review_order"] = cross_final["review_order"]

    cross["header"] = {
        "cells": header_result["cells"],
        "source_document": header_result["source_document"],
    }
    cross["plan_reference"] = plan_reference
    if operator_match:
        # Precedência: a identidade aceite continua a ser a do resolve_operator
        # — as células do cabeçalho descrevem-na, não a substituem.
        cross["operator"] = operator_match

    if edits:
        # Se o humano editou entre o cálculo e a gravação, o CAS recusa e
        # desiste-se — a edição dele dispara um run_cross_check novo.
        db.apply_cross_corrections(conn, uid, data, cross, expected, edits)
    else:
        db.save_cross_check(conn, uid, cross, expected_revision=expected)


# ---------- páginas ----------

@app.get("/", response_class=HTMLResponse)
def home(request: Request, status: str = "", operador: str = "", setor: str = "",
         data: str = "", data_captura: str = "", of: str = "",
         created: str = "", deleted: str = ""):
    conn = _conn()
    try:
        sheets = db.list_sheets(conn, status=status or None, operador=operador or None,
                                setor=setor or None, data_folha=data or None,
                                data_captura=data_captura or None, of=of or None)
        options = db.filter_options(conn)
    finally:
        conn.close()
    # querystring dos filtros não-status, para os chips preservarem os filtros.
    # URL-encoded: um operador «SILVA & VINHA» truncava o filtro no «&».
    from urllib.parse import quote
    parts = [f"&{k}={quote(v)}" for k, v in (("operador", operador), ("setor", setor),
                                             ("data", data), ("data_captura", data_captura),
                                             ("of", of)) if v]
    return templates.TemplateResponse(request, "home.html", {
        "sheets": sheets, "options": options,
        "f": {"status": status, "operador": operador, "setor": setor,
              "data": data, "data_captura": data_captura, "of": of},
        "filter_qs": "".join(parts),
        "created": created, "deleted": deleted,
        "tunnel_url": tunnel_url(),
    })


@app.get("/captura", response_class=HTMLResponse)
def captura(request: Request):
    return templates.TemplateResponse(request, "captura.html", {
        "templates_list": list(TEMPLATES.values()),
    })


@app.get("/captura/camara", response_class=HTMLResponse)
def camara(request: Request):
    return templates.TemplateResponse(request, "camara.html", {
        "templates_list": list(TEMPLATES.values()),
    })


@app.get("/estado", response_class=HTMLResponse)
def estado_page(request: Request, q: str = "", familia: str = "", of: str = ""):
    conn = _conn()
    try:
        sheets = db.list_sheets(conn)
    finally:
        conn.close()
    by_status: dict[str, int] = {}
    for s in sheets:
        by_status[s["status"]] = by_status.get(s["status"], 0) + 1
    return templates.TemplateResponse(request, "estado.html", {
        "estado": estado_data.load_estado(q, familia, of),
        "q": q, "familia": familia, "of": of,
        "by_status": by_status, "n_sheets": len(sheets),
    })


@app.get("/estado/pdf")
def estado_pdf(q: str = "", familia: str = ""):
    data = estado_data.load_estado(q, familia)
    if not data["available"]:
        raise HTTPException(503, f"Postgres indisponível: {data['error']}")
    content = pdf_gen.estado_pdf(data["rows"], q, familia, data["kpis"])
    return Response(content, media_type="application/pdf", headers={
        "Content-Disposition": 'attachment; filename="estado_producao.pdf"',
    })


# rotas antigas — não partir bookmarks
@app.get("/dashboard")
def dashboard_redirect():
    return RedirectResponse("/estado", status_code=301)


@app.get("/capture")
def capture_redirect():
    return RedirectResponse("/captura", status_code=301)


def _save_image(content: bytes, filename: str) -> tuple[str, str]:
    sha = hashlib.sha256(content).hexdigest()
    settings.images_dir.mkdir(parents=True, exist_ok=True)
    dest = settings.images_dir / f"{sha[:16]}_{filename}"
    dest.write_bytes(content)
    return str(dest), sha


def _pdf_to_images(content: bytes, stem: str) -> list[tuple[bytes, str]]:
    """PDF de kanbans digitalizados: cada página vira uma imagem PNG."""
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(content)
    out: list[tuple[bytes, str]] = []
    try:
        for i in range(len(doc)):
            pil = doc[i].render(scale=2.0).to_pil()
            buf = io.BytesIO()
            pil.save(buf, format="PNG")
            out.append((buf.getvalue(), f"{stem}_p{i + 1}.png"))
    finally:
        doc.close()
    return out


# O OCR corre em segundo plano: o túnel Cloudflare corta pedidos aos ~100 s e um
# lote de PDF real (12 páginas × 2 chamadas Gemini) demora minutos. O upload
# responde já; o worker preenche as folhas e o Histórico (auto-refresh) mostra-as.
PROCESS_IN_BACKGROUND = True


def _process_sheet(uid: str, force_ocr: bool = False) -> None:
    conn = db.connect()
    try:
        sheet = db.get_sheet(conn, uid)
        if not sheet or sheet["status"] != "pending" or not sheet.get("image_path"):
            return
        template_name = sheet["template_name"]
        template = get_template(template_name)
        provider = get_provider()
        # O OCR lê a folha na orientação de leitura, não como ela saiu do
        # scanner: com a folha deitada o modelo troca colunas.
        image_path = imaging.render_oriented(
            Path(sheet["image_path"]), int(sheet.get("image_rotation") or 0)
        )
        # Verso em branco do scanner: sem tinta não há nada para ler, e mandar
        # uma página vazia ao modelo produzia folhas inventadas inteiras.
        # `force_ocr` é o revisor a discordar da deteção — respeita-se.
        if not force_ocr and imaging.is_blank_page(image_path, settings.blank_ink_threshold):
            extraction = empty_extraction(template)
            extraction["_blank_page"] = True
            if db.set_extraction(conn, uid, extraction):
                run_cross_check(conn, uid)
            return
        # Folha TPL102 tem frente (produção) e verso (paragens). O provider
        # classifica E transcreve na MESMA chamada — eram duas por página, e a
        # classificação sozinha gastava metade da quota do free tier.
        try:
            if template.family == "cantoneiras" and hasattr(provider, "extract_auto"):
                kinds = {"producao": get_template("cantoneiras_kanban"),
                         "paragens": get_template("cantoneiras_paragens")}
                kind, extraction = provider.extract_auto(image_path, kinds)
                if kinds[kind].name != template_name:
                    # a reclassificação grava-se junto com a transcrição, na
                    # mesma escrita atómica (ver db.set_extraction)
                    template_name = kinds[kind].name
                    template = kinds[kind]
            else:
                extraction = provider.extract(image_path, template)
            # Ponto comum da cadeia (Qwen/Gemini/Claude): se a leitura veio
            # sem identificação no cabeçalho, uma segunda chamada focada na
            # faixa superior tenta recuperá-la. Nunca pisa o que foi lido.
            extraction = rescue_header(provider, image_path, template, extraction)
        except OcrError as exc:
            # OCR falhou (rede, quota, chave): a folha abre vazia para
            # preenchimento manual; o erro fica no trilho de auditoria
            extraction = empty_extraction(template)
            extraction["_ocr_error"] = str(exc)
        if not db.set_extraction(conn, uid, extraction, template_name=template_name):
            return  # o revisor começou a editar entretanto: o trabalho dele manda
        run_cross_check(conn, uid)
    except Exception as exc:  # nunca matar o worker do lote por causa de uma folha
        print(f"[worker] folha {uid}: {exc}", flush=True)
        try:
            err_conn = db.connect()
            try:
                # sem isto a folha ficava `pending` para sempre: spinner
                # eterno na página e nenhum botão para a recuperar
                db.set_error(err_conn, uid, str(exc))
            finally:
                err_conn.close()
        except Exception:
            pass
    finally:
        conn.close()


# Entre folhas: 1 chamada/folha, ~10 RPM — o limite do free tier do primário.
_BATCH_SLEEP_S = 6.0
# Antes da segunda passagem: tempo para um pico de 503 («high demand») passar.
_RETRY_DELAY_S = 60.0
# Só se re-tenta o que recupera sozinho: 5xx e rede. Quota (429) não — essa
# volta pelo botão «Tentar OCR outra vez» ou pelo fallback pago, se existir.
_TRANSIENT_MARKERS = ("HTTP 500", "HTTP 502", "HTTP 503", "HTTP 504", "indisponível")


def _transient_failures(uids: list[str]) -> list[str]:
    """Folhas do lote cujo OCR falhou por causa passageira e ninguém tocou."""
    out: list[str] = []
    conn = db.connect()
    try:
        for uid in uids:
            sheet = db.get_sheet(conn, uid)
            # `in_review` também conta: as escritas do próprio motor (ex.: a
            # data assumida) mudam o estado sem nenhum humano ter tocado — o
            # guarda contra pisar trabalho humano é o edit_count, logo abaixo.
            if not sheet or sheet["status"] not in ("extracted", "in_review"):
                continue                      # pendente/validada/erro: não mexer
            if db.edit_count(conn, uid):
                continue                      # já há trabalho humano em cima
            err = (sheet.get("raw_extraction") or {}).get("_ocr_error") or ""
            if any(m in err for m in _TRANSIENT_MARKERS):
                out.append(uid)
    finally:
        conn.close()
    return out


def _process_batch(uids: list[str]) -> None:
    for i, uid in enumerate(uids):
        _process_sheet(uid)
        if i < len(uids) - 1:
            time.sleep(_BATCH_SLEEP_S)
    # Segunda passagem única pelas falhas passageiras: um pico de 503 a meio de
    # um lote de 26 páginas não deve deixar folhas mortas à espera de cliques.
    retry = _transient_failures(uids)
    if not retry:
        return
    time.sleep(_RETRY_DELAY_S)
    conn = db.connect()
    try:
        for uid in retry:
            db.mark_pending(conn, uid)
    finally:
        conn.close()
    for i, uid in enumerate(retry):
        _process_sheet(uid)
        if i < len(retry) - 1:
            time.sleep(_BATCH_SLEEP_S)


@app.post("/upload")
def upload(template_name: str = Form(...), photos: list[UploadFile] = []):
    # Rota síncrona de propósito: corre no threadpool. Como `async def`, o
    # render do PDF (26 páginas × pypdfium) bloqueava o event loop dezenas de
    # segundos e a app inteira deixava de responder durante um upload.
    try:
        template = get_template(template_name)
    except KeyError:
        raise HTTPException(422, f"Template desconhecido: {template_name}")
    images: list[tuple[bytes, str]] = []
    for up in photos:
        if not up.filename:
            continue
        content = up.file.read()
        name = Path(up.filename).name
        if name.lower().endswith(".pdf") or up.content_type == "application/pdf":
            images.extend(_pdf_to_images(content, Path(name).stem))
        else:
            images.append((content, name))

    conn = _conn()
    try:
        if not images:
            # sem ficheiros: folha vazia para preenchimento manual (imediato)
            uid = db.create_sheet(conn, template_name)
            db.set_extraction(conn, uid, empty_extraction(template))
            run_cross_check(conn, uid)
            uids = [uid]
        else:
            uids = []
            for content, name in images:
                image_path, sha = _save_image(content, name)
                uids.append(db.create_sheet(conn, template_name, image_path, sha))
    finally:
        conn.close()

    if images:
        if PROCESS_IN_BACKGROUND:
            threading.Thread(target=_process_batch, args=(uids,), daemon=True).start()
        else:  # testes: determinístico
            _process_batch(uids)

    if len(uids) == 1:
        return RedirectResponse(f"/sheet/{uids[0]}", status_code=303)
    return RedirectResponse(f"/?created={len(uids)}", status_code=303)


@app.get("/upload")
def upload_get_redirect():
    """Um refresh depois de um upload interrompido não deve dar 405."""
    return RedirectResponse("/captura", status_code=303)


# PDFs de kanban que o scanner da fábrica põe no Drive: «06-08-2026 - Rapid20T 1.pdf»,
# «18-08-2026.PDF»… O padrão vem da config (MES_KANBAN_PDF_RE): o scanner
# mudou a convenção de nomes a 14-08 e o padrão antigo, preso a «rapid»,
# deixava os lotes novos por ingerir.
_KANBAN_PDF_RE = re.compile(settings.kanban_pdf_re, re.IGNORECASE)

# caminhos de célula aceites no /edit; nomes de campo só minúsculas/underscore
_FIELD_PATH_RE = re.compile(
    r"^(?:rows\[(?P<idx>\d{1,4})\]\.(?P<rfield>[a-z_]{1,40})"
    r"|(?P<section>header|footer)\.(?P<sfield>[a-z_]{1,40}))$"
)
_MAX_ROWS = 200   # nenhuma folha física tem 200 linhas


@app.post("/ingest/drive")
def ingest_drive(request: Request):
    """Ingestão dos PDFs de kanban que o sync do Drive deixou no servidor.

    Fecha o ciclo scanner → Drive → app: até aqui o sync trazia os PDFs para a
    máquina e ficavam à espera de um upload manual (4 lotes chegaram a
    acumular-se). Idempotente a dois níveis: sha256 do PDF (versão nova do
    mesmo ficheiro reprocessa) e sha256 de cada página (o que já entrou — por
    upload manual, por exemplo — nunca duplica). O OCR segue em background,
    como no upload manual.
    """
    if settings.admin_token and request.headers.get("X-Admin-Token") != settings.admin_token:
        raise HTTPException(403, "X-Admin-Token inválido.")
    drive = settings.drive_dir
    if not drive.is_dir():
        raise HTTPException(503, f"Pasta do Drive não encontrada: {drive}")

    pdfs = sorted(p for p in drive.iterdir()
                  if p.is_file() and _KANBAN_PDF_RE.match(p.name))
    report = {"pdfs_novos": 0, "folhas_criadas": 0, "paginas_repetidas": 0,
              "pdfs_vistos": len(pdfs)}
    uids: list[str] = []
    conn = _conn()
    try:
        done = db.ingested_shas(conn)
        for pdf in pdfs:
            content = pdf.read_bytes()
            pdf_sha = hashlib.sha256(content).hexdigest()
            if pdf_sha in done:
                continue
            try:
                pages = _pdf_to_images(content, pdf.stem)
            except Exception as exc:
                # PDF estragado não pode encravar o ciclo diário inteiro
                print(f"[ingest] {pdf.name}: PDF ilegível ({exc})", flush=True)
                continue
            report["pdfs_novos"] += 1
            for page_bytes, page_name in pages:
                page_sha = hashlib.sha256(page_bytes).hexdigest()
                if db.image_sha_exists(conn, page_sha):
                    report["paginas_repetidas"] += 1
                    continue
                image_path, sha = _save_image(page_bytes, page_name)
                uids.append(db.create_sheet(conn, "cantoneiras_kanban", image_path, sha))
                report["folhas_criadas"] += 1
            db.record_ingested(conn, pdf.name, pdf_sha, len(pages))
    finally:
        conn.close()

    if uids:
        if PROCESS_IN_BACKGROUND:
            threading.Thread(target=_process_batch, args=(uids,), daemon=True).start()
        else:  # testes: determinístico
            _process_batch(uids)
    return report


@app.post("/sheet/{uid}/reocr")
def sheet_reocr(uid: str, force: int = 0):
    """Re-ler a foto com OCR (ex.: depois de um 429 de quota).

    `?force=1` salta a deteção de página em branco — é o revisor a dizer que a
    página tem mesmo conteúdo, e a palavra dele vale mais que a heurística."""
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
        if not sheet:
            raise HTTPException(404)
        if sheet["status"] == "validated":
            raise HTTPException(409, "Folha validada é imutável.")
        if not sheet.get("image_path"):
            raise HTTPException(422, "Folha sem foto — não há nada para reler.")
        if not force and db.edit_count(conn, uid):
            # já há trabalho humano em cima: re-ler substituía-o todo pela
            # transcrição nova; exige-se a intenção explícita (?force=1)
            raise HTTPException(
                409, "Esta folha já tem correções manuais — re-ler o OCR "
                     "substituía-as. Usa «Forçar OCR» se for mesmo isso que queres.")
        db.mark_pending(conn, uid)
    finally:
        conn.close()
    if PROCESS_IN_BACKGROUND:
        threading.Thread(target=_process_sheet, args=(uid, bool(force)), daemon=True).start()
    else:
        _process_sheet(uid, bool(force))
    return RedirectResponse(f"/sheet/{uid}", status_code=303)


@app.post("/sheet/{uid}/delete")
def sheet_delete(uid: str):
    conn = _conn()
    try:
        try:
            image_path = db.delete_sheet(conn, uid)
        except KeyError:
            raise HTTPException(404)
        except PermissionError:
            raise HTTPException(409, "Folha validada é imutável — não se apaga.")
        # O ficheiro pode ser partilhado (frente/verso, uploads repetidos):
        # só se apaga quando NENHUMA outra folha ainda aponta para ele.
        shared = bool(image_path) and db.image_path_in_use(conn, image_path)
    finally:
        conn.close()
    if image_path and not shared:
        p = Path(image_path).resolve()
        if p.is_relative_to(settings.images_dir.resolve()) and p.is_file():
            imaging.clear_renders(p)   # os .rotN.png derivados vão junto
            p.unlink()
    return RedirectResponse("/?deleted=1", status_code=303)


@app.get("/sheet/{uid}/csv")
def sheet_csv(uid: str):
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
    finally:
        conn.close()
    if not sheet:
        raise HTTPException(404)
    template = get_template(sheet["template_name"])
    data = sheet["sheet_data"] or {}
    header = data.get("header") or {}
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["folha", "estado", "operador", "data", "setor_maquina", "linha"]
               + list(template.row_fields))
    for i, row in enumerate(data.get("rows") or []):
        if not any(v is not None and str(v).strip() for v in row.values()):
            continue
        w.writerow([
            sheet["uid"][:8], sheet["status"], header.get("operador"),
            header.get("data"), header.get("setor_maquina"), i + 1,
        ] + [row.get(f) for f in template.row_fields])
    return Response(buf.getvalue(), media_type="text/csv; charset=utf-8", headers={
        "Content-Disposition": f'attachment; filename="kanban_{uid[:8]}.csv"',
    })


@app.get("/export/cpis")
def export_cpis(de: str = "", ate: str = "", operador: str = "", validadas: int = 0):
    """A tabela plana da Metalogalva 2 («MigracaoNikufraCPIS_….xlsx»).

    Fonte: o staging local (todas as folhas de produção com data) — funciona
    com o Postgres em baixo e cobre também o que ainda está em revisão;
    `?validadas=1` restringe ao que já foi validado. Filtros `de`/`ate` em
    ISO (aaaa-mm-dd) e `operador` por nome.
    """
    conn = _conn()
    try:
        sheets = db.list_sheets(conn, status="validated" if validadas else None)
        cpis_rows: list[dict] = []
        for meta in sheets:
            if "paragens" in meta["template_name"]:
                continue
            sheet = db.get_sheet(conn, meta["uid"])
            if not sheet or not sheet.get("sheet_data"):
                continue
            header = sheet["sheet_data"].get("header") or {}
            if operador and str(header.get("operador") or "").strip() != operador:
                continue
            try:
                iso = pg_store.normalize_sheet_date(header.get("data"))
            except pg_store.InvalidSheetDate:
                iso = None
            if de and (not iso or iso < de):
                continue
            if ate and (not iso or iso > ate):
                continue
            cross = sheet.get("cross_check") or {}
            cross_rows = {r["row_index"]: r for r in cross.get("rows", [])}
            op = cross.get("operator")
            for i, row in enumerate(sheet["sheet_data"].get("rows") or []):
                if not any(v is not None and str(v).strip() for v in row.values()):
                    continue
                cpis_rows.append(
                    (iso or "9999", str(header.get("operador") or ""), meta["uid"], i,
                     cpis_export.cpis_row_for(sheet, row, cross_rows.get(i), op)))
    finally:
        conn.close()
    # ordenação do original: data, operador, folha, linha
    cpis_rows.sort(key=lambda t: t[:4])
    content = cpis_export.build_cpis_workbook([t[4] for t in cpis_rows])
    filename = cpis_export.cpis_filename_for(de or None, ate or None, bool(validadas))
    return Response(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/export/basedados")
def export_basedados(de: str = "", ate: str = "", operador: str = "", validadas: int = 0):
    """A tabela plana no formato Modelo_BaseDados_PerfisCantoneiras.xlsx
    (11 colunas), comum aos dois setores kanban. Filtros de período/operador
    como no CPIS, mas SÓ folhas validadas: a BaseDados é o registo oficial e
    uma folha por rever ainda pode mudar. O query param `validadas` continua a
    ser aceite (links/bookmarks antigos), mas é ignorado."""
    del validadas
    conn = _conn()
    try:
        sheets = db.list_sheets(conn, status="validated")
        bd_rows: list[tuple] = []
        for meta in sheets:
            if "paragens" in meta["template_name"]:
                continue
            sheet = db.get_sheet(conn, meta["uid"])
            if not sheet or not sheet.get("sheet_data"):
                continue
            header = sheet["sheet_data"].get("header") or {}
            if operador and str(header.get("operador") or "").strip() != operador:
                continue
            try:
                iso = pg_store.normalize_sheet_date(header.get("data"))
            except pg_store.InvalidSheetDate:
                iso = None
            if de and (not iso or iso < de):
                continue
            if ate and (not iso or iso > ate):
                continue
            cross = sheet.get("cross_check") or {}
            cross_rows = {r["row_index"]: r for r in cross.get("rows", [])}
            op = cross.get("operator")
            for i, row in enumerate(sheet["sheet_data"].get("rows") or []):
                if not any(v is not None and str(v).strip() for v in row.values()):
                    continue
                bd_rows.append(
                    (iso or "9999", str(header.get("operador") or ""), meta["uid"], i,
                     cpis_export.basedados_row_for(sheet, row, cross_rows.get(i), op)))
    finally:
        conn.close()
    bd_rows.sort(key=lambda t: t[:4])
    content = cpis_export.build_basedados_workbook([t[4] for t in bd_rows])
    filename = cpis_export.basedados_filename_for(de or None, ate or None)
    return Response(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/export.xlsx")
def export_xlsx():
    """Todas as linhas de produção validadas (Postgres) num Excel."""
    import openpyxl
    import psycopg

    try:
        with psycopg.connect(pg_store._dsn(), connect_timeout=5) as pconn:
            pconn.read_only = True
            with pconn.cursor() as cur:
                cur.execute(
                    "SELECT sheet_uid, row_index, sheet_date, family, operator_name, "
                    "machine, production_order, sales_order, customer_name, model_ref, "
                    "matched_plan_key, match_confidence, quantity, length_mm, width_mm, "
                    "thickness_mm, hours_worked, validated_at "
                    "FROM mes_kanban.production_records "
                    "ORDER BY sheet_date DESC, sheet_uid, row_index")
                cols = [d.name for d in cur.description]
                rows = cur.fetchall()
    except Exception as exc:
        raise HTTPException(503, f"Postgres indisponível: {exc}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "producao"
    ws.append(cols)
    for r in rows:
        ws.append([str(v) if v is not None and not isinstance(v, (int, float)) else v
                   for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="producao_mes.xlsx"'})


def _safe_back(back: str | None) -> str | None:
    """Só aceita caminhos internos — impede que um ?back= leve para fora do site.
    O `\\` conta como `//`: os browsers normalizam `/\\evil.com` para
    `//evil.com` e o filtro de prefixo deixava-o passar."""
    if not back or not back.startswith("/") or back.startswith("//") or "\\" in back:
        return None
    return back


def _diverged_map(sheet: dict) -> dict[str, str]:
    """Células onde o valor atual já não é o que a máquina leu.

    Devolve {field_path: valor_lido_pelo_ocr} — o valor original serve de
    tooltip, para o revisor saber de onde é que a célula veio sem ter de
    trocar de vista. Cobre cabeçalho, linhas e rodapé.
    """
    raw = sheet.get("raw_extraction") or {}
    cur = sheet.get("sheet_data") or {}
    if not raw:
        return {}
    out: dict[str, str] = {}

    def diff(path: str, a, b) -> None:
        sa, sb = str(a or "").strip(), str(b or "").strip()
        if sa != sb:
            out[path] = sa

    for section in ("header", "footer"):
        raw_sec, cur_sec = raw.get(section) or {}, cur.get(section) or {}
        for f in set(raw_sec) | set(cur_sec):
            diff(f"{section}.{f}", raw_sec.get(f), cur_sec.get(f))

    raw_rows, cur_rows = raw.get("rows") or [], cur.get("rows") or []
    for i in range(max(len(raw_rows), len(cur_rows))):
        r = raw_rows[i] if i < len(raw_rows) and isinstance(raw_rows[i], dict) else {}
        c = cur_rows[i] if i < len(cur_rows) and isinstance(cur_rows[i], dict) else {}
        for f in set(r) | set(c):
            diff(f"rows[{i}].{f}", r.get(f), c.get(f))
    return out


@app.get("/sheet/{uid}", response_class=HTMLResponse)
def sheet_view(request: Request, uid: str, back: str | None = None,
               view: str | None = None):
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
    finally:
        conn.close()
    if not sheet:
        raise HTTPException(404)
    template = get_template(sheet["template_name"])
    raw = sheet.get("raw_extraction") or {}
    raw_rows = [r for r in (raw.get("rows") or []) if isinstance(r, dict)]
    has_ocr = bool(sheet.get("image_path")) and any(
        v is not None and str(v).strip() for r in raw_rows for v in r.values()
    )
    # Vista crua: mostra a transcrição original e DESLIGA as cores. As cores do
    # cross-check validam o valor final contra o plano — pintá-las por cima de
    # valores crus seria dizer que o motor aprovou o que ele nunca viu.
    view_mode = "raw" if (view == "raw" and raw) else "final"
    cross_rows = {}
    if view_mode == "final" and sheet["cross_check"]:
        cross_rows = {
            r["row_index"]: {**r, "cells_by_field": {c["field"]: c for c in r["cells"]}}
            for r in sheet["cross_check"]["rows"]
        }
    diverged = _diverged_map(sheet) if view_mode == "final" else {}
    cross = sheet["cross_check"] or {}
    header_cross_data = cross.get("header") or {}
    header_cells = header_cross_data.get("cells") or {}
    source_document = (
        header_cross_data.get("source_document") or _source_document(sheet)
    )
    return templates.TemplateResponse(request, "sheet.html", {
        "sheet": sheet, "t": template, "cross_rows": cross_rows,
        "summary": cross.get("summary"),
        "review_order": cross.get("review_order", []),
        "plan_reference": cross.get("plan_reference") or {},
        "stored": request.query_params.get("stored"),
        "has_ocr": has_ocr, "view_mode": view_mode,
        "operator": cross.get("operator"),
        "header_cells": header_cells,
        "header_labels": header_cross.HEADER_LABELS,
        "source_document": source_document,
        "diverged": diverged, "n_diverged": len(diverged),
        "back_url": _safe_back(back),
    })


@app.get("/sheet/{uid}/photo")
def sheet_photo(uid: str, original: int = 0):
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
    finally:
        conn.close()
    if not sheet or not sheet.get("image_path"):
        raise HTTPException(404)
    path = Path(sheet["image_path"]).resolve()
    # a foto tem de viver dentro da pasta de imagens da app (anti path-traversal)
    if not path.is_relative_to(settings.images_dir.resolve()) or not path.is_file():
        raise HTTPException(404)
    if not original:
        path = imaging.render_oriented(path, int(sheet.get("image_rotation") or 0))
    return FileResponse(path)


@app.post("/sheet/{uid}/rotate")
def sheet_rotate(uid: str):
    """Roda mais 90° no sentido horário, por cima da correcção automática.

    A automática acerta em todas as digitalizações que vimos, mas é um palpite
    sobre o conteúdo a partir da forma da imagem — se sair ao contrário, o
    revisor resolve com cliques.
    """
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
        if not sheet:
            raise HTTPException(404)
        rotation = db.set_image_rotation(conn, uid, int(sheet.get("image_rotation") or 0) + 90)
    finally:
        conn.close()
    return {"ok": True, "rotation": rotation}


def _totais_plano(linhas: list) -> dict:
    """Somas de planeada/feita/falta das linhas do pop-up.

    Cada soma só existe se houver pelo menos um valor na coluna — somar zeros
    de colunas vazias mostraria «0» onde a verdade é «não se sabe».
    """
    totais: dict = {"planeada": None, "feita": None, "falta": None,
                    "parcial": len(linhas) >= 500}
    colunas = {"planeada": "quantity_planned", "feita": "quantity_made",
               "falta": "remaining_quantity"}
    for chave, col in colunas.items():
        valores = [l[col] for l in linhas if l.get(col) is not None]
        if valores:
            totais[chave] = float(sum(valores))
    return totais


@app.get("/sheet/{uid}/plano/{row_index}", response_class=HTMLResponse)
def sheet_plano_perfil(request: Request, uid: str, row_index: int, origem: str = ""):
    """As referências do plano para a chave OF + Perfil de uma linha.

    Recebe a linha e não a chave: é o servidor que resolve a OF (incluindo a
    herdada da linha de cima) e a forma canónica do perfil. Se fosse o template
    a montar `?of=&perfil=`, teria de conhecer as convenções do plano e podia
    perguntar por uma chave diferente daquela com que o motor cruzou.

    `origem=perf_comp` = o clique veio da marca de perfil completo, que afirma
    «fiz a quantidade toda»: o pop-up junta os totais e avisa se o plano ainda
    mostra falta.
    """
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
        human = db.human_fields_by_row(conn, uid)
    finally:
        conn.close()
    if not sheet:
        raise HTTPException(404)
    template = get_template(sheet["template_name"])
    rows = (sheet["sheet_data"] or {}).get("rows") or []
    if row_index < 0 or row_index >= len(rows):
        raise HTTPException(404)

    ctx: dict = {"row_index": row_index, "of": None, "perfil": None,
                 "linhas": [], "perfis": [], "erro": None, "plano": {},
                 "origem_perf_comp": origem == "perf_comp", "totais": None}
    try:
        index = get_index(template.index_loader) if template.index_loader else None
        if index is None:
            ctx["erro"] = "Esta folha não cruza com o plano."
            return templates.TemplateResponse(request, "_plano_perfil.html", ctx)

        content = tuple(f.name for f in index.spec.identity_fields
                        if f.name not in carryover.CARRY_FIELDS)
        identities = carryover.resolve(rows, content, human)
        eff = carryover.effective_row(rows[row_index], identities[row_index])
        escrito = str(eff.get("perfil") or "").strip()
        of_escrita = str(eff.get("of") or "").strip()
        # A OF no plano leva prefixo; procurar pela forma que lá existe.
        of = next(
            (index.entries[i]["of"] for i in index.exact_matches("of", of_escrita)),
            of_escrita,
        ) if of_escrita else ""
        perfil = index.normalize_written("perfil", escrito) if escrito else ""
        ctx.update({"of": of, "perfil": perfil, "perfil_escrito": escrito,
                    "herdou_of": identities[row_index].is_inherited("of")})
        if not of:
            ctx["erro"] = "Esta linha não tem OF — escreve-a (ou herda-a da linha de cima)."
        else:
            ctx["plano"] = loaders.plan_snapshot_info()
            if perfil:
                ctx["linhas"] = loaders.fetch_profile_lines(of, perfil)
            if ctx["linhas"]:
                ctx["totais"] = _totais_plano(ctx["linhas"])
            else:
                # Sem correspondência mostra-se o que a obra tem mesmo: o caso
                # comum é o perfil estar escrito com uma medida trocada.
                ctx["perfis"] = loaders.fetch_profiles_in_of(of)
    except Exception as exc:  # Postgres em baixo não pode rebentar a revisão
        ctx["erro"] = f"Não foi possível ler o plano: {exc}"
    return templates.TemplateResponse(request, "_plano_perfil.html", ctx)


@app.get("/sheet/{uid}/pdf")
def sheet_pdf(uid: str):
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
        edits = db.edit_count(conn, uid)
    finally:
        conn.close()
    if not sheet:
        raise HTTPException(404)
    template = get_template(sheet["template_name"])
    content = pdf_gen.sheet_pdf(sheet, template, edits)
    return Response(content, media_type="application/pdf", headers={
        "Content-Disposition": f'attachment; filename="kanban_{uid[:8]}.pdf"',
    })


@app.post("/sheet/{uid}/edit")
def sheet_edit(uid: str, field_path: str = Form(...), value: str = Form(""),
               revision: int = Form(...), actor: str = Form("operador")):
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
        if not sheet:
            raise HTTPException(404)
        if sheet["status"] == "validated":
            raise HTTPException(409, "Folha validada é imutável.")
        data = sheet["sheet_data"] or {"header": {}, "rows": [], "footer": {}}
        # field_path: 'header.data' | 'rows[3].of' | 'footer.horas_trabalhadas'.
        # Validação estrita: o endpoint é público (túnel sem auth) e um
        # rows[50000000] criava 50 M de dicts — OOM e queda do processo; um
        # índice negativo corrompia a proteção de células humanas.
        m = _FIELD_PATH_RE.match(field_path)
        if not m:
            raise HTTPException(400, "field_path inválido.")
        old = None
        value_clean = value.strip() or None
        if m.group("idx") is not None:
            i = int(m.group("idx"))
            if i > _MAX_ROWS:
                raise HTTPException(422, f"Linha {i} fora do limite ({_MAX_ROWS}).")
            fname = m.group("rfield")
            while len(data["rows"]) <= i:
                data["rows"].append({})
            old = data["rows"][i].get(fname)
            data["rows"][i][fname] = value_clean
        else:
            section, fname = m.group("section"), m.group("sfield")
            old = (data.get(section) or {}).get(fname)
            data.setdefault(section, {})[fname] = value_clean
        # Edição que não muda nada (ex.: focar a célula e sair) não é uma
        # decisão humana: gravá-la marcava o campo como inviolável e
        # desligava a herança sem o revisor querer.
        old_clean = str(old).strip() or None if old is not None else None
        if old_clean == value_clean:
            return RedirectResponse(f"/sheet/{uid}", status_code=303)
        # controlo otimista: a revisão vem do formulário — se a folha mudou
        # desde que a página foi carregada, recusa em vez de sobrescrever
        if not db.save_sheet_data(conn, uid, data, revision):
            raise HTTPException(409, "A folha mudou entretanto — recarrega a página.")
        db.record_edit(conn, uid, field_path, old, value_clean, "human", actor)
        run_cross_check(conn, uid)
    finally:
        conn.close()
    return RedirectResponse(f"/sheet/{uid}", status_code=303)


@app.post("/sheet/{uid}/add-row")
def add_row(uid: str):
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
        if not sheet:
            raise HTTPException(404)
        if sheet["status"] == "validated":
            raise HTTPException(409, "Folha validada é imutável.")
        data = sheet["sheet_data"]
        if data is None:
            raise HTTPException(409, "A folha ainda está a ser lida pelo OCR.")
        template = get_template(sheet["template_name"])
        data["rows"].append({f: None for f in template.row_fields})
        if not db.save_sheet_data(conn, uid, data, sheet["revision"]):
            raise HTTPException(409, "A folha mudou entretanto — recarrega a página.")
    finally:
        conn.close()
    return RedirectResponse(f"/sheet/{uid}", status_code=303)


@app.post("/sheet/{uid}/recheck")
def recheck(uid: str):
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
        if not sheet:
            raise HTTPException(404)
        if sheet["status"] == "validated":
            raise HTTPException(409, "Folha validada é imutável.")
        run_cross_check(conn, uid)
    finally:
        conn.close()
    return RedirectResponse(f"/sheet/{uid}", status_code=303)


@app.post("/sheet/{uid}/validate")
def validate(uid: str, actor: str = Form("operador")):
    """A única porta para o Postgres: valida → INSERT em mes_kanban → imutável."""
    conn = _conn()
    try:
        sheet = db.get_sheet(conn, uid)
        if not sheet:
            raise HTTPException(404)
        if sheet["status"] == "validated":
            raise HTTPException(409, "Folha já validada.")
        header = (sheet["sheet_data"] or {}).get("header") or {}
        if not str(header.get("operador") or "").strip():
            raise HTTPException(422, "Validação exige operador preenchido no cabeçalho.")
        if not str(header.get("data") or "").strip():
            raise HTTPException(422, "Validação exige data preenchida no cabeçalho.")
        template = get_template(sheet["template_name"])
        try:
            n = pg_store.store_validated_sheet(
                sheet, template, db.edit_count(conn, uid), actor)
        except pg_store.InvalidSheetDate as exc:
            raise HTTPException(
                422, f"Data «{exc}» não é interpretável — escreve dd/mm/aaaa.")
        db.mark_validated(conn, uid, actor)
    finally:
        conn.close()
    return RedirectResponse(f"/sheet/{uid}?stored={n}", status_code=303)
