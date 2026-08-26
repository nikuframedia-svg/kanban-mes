"""Export CPIS — a tabela plana da Metalogalva 2, célula por célula.

Réplica do `/export/cpis` do sistema original (nikuframedia-svg/ocr,
backend/app/web/export.py): folha única «Folha1», 26 colunas com cabeçalhos de
negócio pela mesma ordem, tipos nativos, number_format por coluna, larguras
fixas, freeze do cabeçalho e neutralização anti-fórmula. O nome do ficheiro
(«MigracaoNikufraCPIS_…») e o layout mantêm-se — é este formato que os
sistemas a jusante da Metalogalva consomem.

As colunas de chapa/pesos que as cantoneiras não têm (M², Nesting, CONI,
pesos…) seguem vazias: o CONTRATO do ficheiro é o mesmo, os dados são os que
esta fábrica produz.
"""

from __future__ import annotations

import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from ..matching import similarity as sim
from ..pg_store import InvalidSheetDate, normalize_sheet_date

# (chave, cabeçalho) — ordem EXATA do original
CPIS_COLUMNS: tuple[tuple[str, str], ...] = (
    ("data", "Data"),
    ("cod_funcionario", "Cód. Funcionário"),
    ("nome_funcionario", "Nome Funcionário"),
    ("setor_maquina_desc", "Setor / Máquina Desc."),
    ("cod_maquina", "Cód. Máquina"),
    ("of", "OF"),
    ("ov", "OV"),
    ("cliente", "Cliente"),
    ("modelo", "Modelo"),
    ("qtd", "QTD"),
    ("qtd_metros", "Qtd Metros"),
    ("m2", "M²"),
    ("nesting", "Nesting"),
    ("cesta_n", "Cesta Nº"),
    ("duracao", "Duração"),
    ("comp_mm", "Comprimento (mm)"),
    ("larg_mm", "Largura (mm)"),
    ("esp_mm", "Espessura (mm)"),
    ("coni", "CONI"),
    ("n_chapas", "Nº Chapas"),
    ("peso_consumido_t", "Peso Consumido (t)"),
    ("peso_produzido_t", "Peso Produzido (t)"),
    ("desperdicio_t", "Desperdício (t)"),
    ("desperdicio_pct", "% Desperdício"),
    ("sucata", "Sucata"),
    ("lote", "Lote"),
)

# larguras do original, uma por coluna
_WIDTHS = (12, 16, 24, 24, 12, 10, 10, 20, 24, 8, 12, 10, 14, 10, 10,
           16, 14, 12, 10, 10, 16, 16, 14, 12, 10, 14)

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FILL_HEADER = PatternFill("solid", start_color="2F5597")
_FONT_BASE = Font(name="Inter", size=10)
_FONT_HEADER = Font(name="Inter", size=11, bold=True, color="FFFFFF")


def neutralize_xlsx(value: object) -> object:
    """Anti-injeção de fórmulas: um valor de OCR/plano que comece por =, +, -
    ou @ seria executado pelo Excel como fórmula ao abrir o ficheiro."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _num(value: object) -> float | None:
    return sim.parse_number(value)


def cpis_row_for(sheet: dict, row: dict, cr: dict | None,
                 operator: dict | None) -> dict:
    """Uma linha de kanban → uma linha CPIS.

    Depois da política de substituição, os valores da linha JÁ são os do
    plano; a herança e as propostas do cross servem de fallback para folhas
    antigas ou por rever.
    """
    header = (sheet.get("sheet_data") or {}).get("header") or {}
    cells = {c["field"]: c for c in (cr or {}).get("cells", [])}

    def efetivo(field: str) -> str:
        value = str(row.get(field) or "").strip()
        cell = cells.get(field) or {}
        if not value:
            value = str(cell.get("inherited") or "").strip()
        if not value:
            value = str(cell.get("proposal") or "").strip()
        return value

    try:
        data = dt.date.fromisoformat(normalize_sheet_date(header.get("data")))
    except (InvalidSheetDate, ValueError):
        data = str(header.get("data") or "").strip() or None

    op = operator or {}
    return {
        "data": data,
        "cod_funcionario": op.get("cod") or str(header.get("n_operador") or "").strip() or None,
        "nome_funcionario": op.get("name") or str(header.get("operador") or "").strip() or None,
        "setor_maquina_desc": str(header.get("setor_maquina") or "").strip() or None,
        "cod_maquina": None,
        "of": sim.strip_ref_prefix(efetivo("of")) or None,
        "ov": sim.strip_ref_prefix(efetivo("ov")) or None,
        "cliente": efetivo("cliente") or None,
        "modelo": efetivo("modelo") or str(row.get("perfil") or "").strip() or None,
        "qtd": _num(row.get("qtd") or row.get("repeticoes")),
        "qtd_metros": (cr or {}).get("line_meters"),
        "m2": None,
        "nesting": str(row.get("nesting") or "").strip() or None,
        "cesta_n": None,
        "duracao": None,
        "comp_mm": (cr or {}).get("plan_length_mm") or _num(row.get("comp_mm")),
        "larg_mm": _num(row.get("larg_mm")),
        "esp_mm": _num(row.get("esp")),
        "coni": None,
        "n_chapas": None,
        "peso_consumido_t": None,
        "peso_produzido_t": None,
        "desperdicio_t": None,
        "desperdicio_pct": None,
        "sucata": _num(row.get("sucata")),
        "lote": str(row.get("lote") or "").strip() or None,
    }


def build_cpis_workbook(cpis_rows: list[dict]) -> bytes:
    """A folha «Folha1» com o layout exato do original."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Folha1"

    for ci, (_, label) in enumerate(CPIS_COLUMNS, start=1):
        cell = ws.cell(row=1, column=ci, value=label)
        cell.font = _FONT_HEADER
        cell.fill = _FILL_HEADER
        cell.border = _BORDER

    for ri, cpis in enumerate(cpis_rows, start=2):
        for ci, (key, _) in enumerate(CPIS_COLUMNS, start=1):
            value = neutralize_xlsx(cpis.get(key))
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.font = _FONT_BASE
            cell.border = _BORDER
            if key == "data" and isinstance(value, dt.date):
                cell.number_format = "DD-MM-YYYY"
            elif key in ("peso_consumido_t", "peso_produzido_t", "desperdicio_t"):
                cell.number_format = "0.000"
            elif key == "desperdicio_pct":
                cell.number_format = "0.00"
            elif key == "esp_mm":
                cell.number_format = "0.0"
            elif key in ("n_chapas", "sucata"):
                cell.number_format = "0"

    for ci, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = width
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def cpis_filename_for(de: str | None, ate: str | None, validadas: bool) -> str:
    """«MigracaoNikufraCPIS_{período}[_validadas].xlsx», como no original."""
    if de and ate and de == ate:
        periodo = f"1-dia_{de}"
    elif de or ate:
        periodo = f"{de or 'inicio'}_{ate or 'hoje'}"
    else:
        periodo = "sempre"
    sufixo = "_validadas" if validadas else ""
    return f"MigracaoNikufraCPIS_{periodo}{sufixo}.xlsx"


# ---------------------------------------------------------------------------
# Export «BaseDados» — o formato Modelo_BaseDados_PerfisCantoneiras.xlsx
# pedido pelo Luís (19/08): folha única «Folha1», 11 colunas, igual nos dois
# setores (cantoneiras MTG3 e perfis MTG2). Cod. Maquina segue vazio até
# existir a tabela oficial de códigos SAP das máquinas.
# ---------------------------------------------------------------------------

# (chave, cabeçalho) — ordem e grafia EXATAS do ficheiro modelo
BASEDADOS_COLUMNS: tuple[tuple[str, str], ...] = (
    ("data", "Data"),
    ("operador_id", "OperadorID"),
    ("nome_operador", "Nome Operador"),
    ("cod_maquina", "Cod. Maquina"),
    ("maquina", "Maquina"),
    ("ov", "OV"),
    ("of", "OF"),
    ("perfil", "Perfil"),
    ("modelo", "Modelo"),
    ("qtd_un", "Qtd [un.]"),
    ("qtd_m", "Qtd [m]"),
)

_BASEDADOS_WIDTHS = (12, 12, 24, 12, 18, 12, 12, 16, 18, 10, 10)


def _efetivo(row: dict, cells: dict, field: str) -> str:
    """Valor efetivo de uma célula: o escrito, senão o herdado.

    Sem fallback para a proposta do cross: com a substituição total, o que o
    motor aprova JÁ está gravado na folha — exportar uma proposta não aplicada
    seria exportar o que a folha não diz."""
    value = str(row.get(field) or "").strip()
    if not value:
        value = str((cells.get(field) or {}).get("inherited") or "").strip()
    return value


def basedados_row_for(sheet: dict, row: dict, cr: dict | None,
                      operator: dict | None) -> dict:
    """Uma linha de kanban → uma linha BaseDados (11 colunas)."""
    header = (sheet.get("sheet_data") or {}).get("header") or {}
    cells = {c["field"]: c for c in (cr or {}).get("cells", [])}

    try:
        data = dt.date.fromisoformat(normalize_sheet_date(header.get("data")))
    except (InvalidSheetDate, ValueError):
        data = str(header.get("data") or "").strip() or None

    op = operator or {}
    return {
        "data": data,
        "operador_id": op.get("cod") or str(header.get("n_operador") or "").strip() or None,
        "nome_operador": op.get("name") or str(header.get("operador") or "").strip() or None,
        "cod_maquina": None,
        "maquina": str(header.get("setor_maquina") or "").strip() or None,
        "ov": sim.strip_ref_prefix(_efetivo(row, cells, "ov")) or None,
        "of": sim.strip_ref_prefix(_efetivo(row, cells, "of")) or None,
        "perfil": _efetivo(row, cells, "perfil") or None,
        "modelo": _efetivo(row, cells, "modelo") or None,
        "qtd_un": _num(row.get("qtd") or row.get("repeticoes")),
        "qtd_m": (cr or {}).get("line_meters"),
    }


def build_basedados_workbook(rows: list[dict]) -> bytes:
    """A folha «Folha1» com o cabeçalho exato do ficheiro modelo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Folha1"

    for ci, (_, label) in enumerate(BASEDADOS_COLUMNS, start=1):
        cell = ws.cell(row=1, column=ci, value=label)
        cell.font = _FONT_HEADER
        cell.fill = _FILL_HEADER
        cell.border = _BORDER

    for ri, bd in enumerate(rows, start=2):
        for ci, (key, _) in enumerate(BASEDADOS_COLUMNS, start=1):
            value = neutralize_xlsx(bd.get(key))
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.font = _FONT_BASE
            cell.border = _BORDER
            if key == "data" and isinstance(value, dt.date):
                cell.number_format = "DD-MM-YYYY"
            elif key == "qtd_m":
                cell.number_format = "0.00"

    for ci, width in enumerate(_BASEDADOS_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = width
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def basedados_filename_for(de: str | None, ate: str | None) -> str:
    """Sem sufixo condicional: a BaseDados exporta SEMPRE só validadas."""
    if de and ate and de == ate:
        periodo = f"1-dia_{de}"
    elif de or ate:
        periodo = f"{de or 'inicio'}_{ate or 'hoje'}"
    else:
        periodo = "sempre"
    return f"BaseDados_{periodo}.xlsx"
