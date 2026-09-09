"""Review/archival contracts, with no writes to a real Postgres database."""
from copy import deepcopy
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

import pytest
from openpyxl import load_workbook

from app import db
from app.matching import loaders
from app.matching.full_profile import attach_plan_facts, expand_entries, plan_identity
from app.matching.refs import PlanIndex
from app.templates_spec import get_template
from app.web import export_routes, export_source, main, plan_review
from tests.test_web import client, create_sheet, edit, get_revision  # noqa: F401

MTG2 = plan_review.IS_MTG2
TEMPLATE = "tpl999_kanban" if MTG2 else "cantoneiras_kanban"
PROFILE = "60.3x2.9" if MTG2 else "100x10"
OTHER_PROFILE = "60.3x3.2" if MTG2 else "100x12"


def entries():
    return [{
        "plan_key": f"old:{i}", "snapshot_id": "old", "production_order_no": "OF42",
        "sales_order_no": "OV21", "customer_name": "CLIENTE", "component_ref": f"REF-{i}",
        "profile_type": PROFILE, "profile_excel_o": "60,3", "material_description": f"DESIGNAÇÃO {i}",
        "length_mm": (1000 + i * 500), "quantity_planned": (8 + i), "quantity_made": (8 + i - qty),
        "remaining_quantity": qty, "remaining_valid": True, "remaining_rule": "Qtd em Falta" if MTG2 else "QTD-Maq.",
        "closed_x": i != 0, "overproduction_quantity": 0,
    } for i, qty in enumerate((5, 2, 0))]


def sheet(full=True):
    cross = expand_entries(entries(), "old", precision=3 if MTG2 else 2) if full else {
        "matched_plan_key": "old:0", "plan_identity": plan_identity(entries()[0]),
        "plan_length_mm": 1000, "plan_line_meters": 4, "line_meters": 4,
    }
    return {"uid": "archive-only", "sheet_no": 77, "revision": 2, "status": "validated", "template_name": TEMPLATE,
            "sheet_data": {"header": {"operador": "ANA", "n_operador": "10", "setor_maquina": "Vanguard" if MTG2 else "FICEP",
                                      "data": "2026-09-01", "turno": "M"},
                           "rows": [{"of": "42", "ov": "21", "cliente": "CLIENTE", "perfil": PROFILE,
                                     "modelo": None if full else "REF-0", "qtd": "4", "perf_comp": "X" if full else None}],
                           "footer": {}},
            "cross_check": {"snapshot_id": "old", "rows": [{"row_index": 0, **cross}]}}


def insert_local(source):
    conn = db.connect()
    try:
        uid = db.create_sheet(conn, source["template_name"])
        db.set_extraction(conn, uid, source["sheet_data"])
        current = db.get_sheet(conn, uid)
        db.save_cross_check(conn, uid, source["cross_check"], current["revision"])
        if source["status"] == "validated":
            conn.execute("UPDATE sheets SET status='validated' WHERE uid=?", (uid,))
            conn.commit()
        return uid
    finally:
        conn.close()


def test_complete_profile_freezes_5_2_0_and_exports_only_positive_children(monkeypatch):
    source = sheet()
    monkeypatch.setattr(loaders, "_fetch", lambda *_: pytest.fail("frozen facts must work offline"))
    prepared = export_source.prepare_sheets([source])[0]
    refs = prepared["cross_check"]["rows"][0]["plan_refs"]
    assert [r["assumed_quantity"] for r in refs] == [5, 2, 0]
    ctx = plan_review.context(prepared, 0, get_template(TEMPLATE), "/")
    assert [r["made_in_sheet"] for r in ctx["linhas"]] == [5, 2, 0]
    assert [r["remaining_after"] for r in ctx["linhas"]] == [0, 0, 0]
    assert ctx["totais"]["nesta_folha"] == 7
    assert ctx["plano"]["frozen"]
    for kind in ("basedados", "cpis"):
        ws = load_workbook(BytesIO(export_routes.workbook(kind, [prepared]))).active
        headers = [c.value for c in ws[1]]
        assert ws.max_row == 3
        qty = headers.index("Qtd [un.]" if kind == "basedados" else "QTD") + 1
        model = headers.index("Modelo") + 1
        assert [ws.cell(i, qty).value for i in (2, 3)] == [5, 2]
        assert [ws.cell(i, model).value for i in (2, 3)] == ["REF-0", "REF-1"]
        if MTG2:
            profile = headers.index("Perfil") + 1
            assert [ws.cell(i, profile).value for i in (2, 3)] == ["60,3", "60,3"]
            if kind == "basedados":
                assert ws.max_column == 12
                assert [ws.cell(i, 12).value for i in (2, 3)] == ["DESIGNAÇÃO 0", "DESIGNAÇÃO 1"]
        elif kind == "basedados":
            assert ws.max_column == 11
    technical = load_workbook(BytesIO(export_routes.technical_workbook([prepared]))).active
    assert technical.max_row == 3
    assert source == sheet(), "preparation must not mutate archived JSON"


def test_full_profile_uses_thickness_and_stable_keys():
    items = entries() + [{**entries()[0], "plan_key": "other-wall", "profile_type": OTHER_PROFILE}]
    indexed = [{**r, "of": r["production_order_no"], "perfil": r["profile_type"], "modelo": r["component_ref"],
                "ov": r["sales_order_no"], "comp_mm": r["length_mm"]} for r in items]
    spec = loaders.PERFIS_SPEC if MTG2 else loaders.CANTONEIRAS_SPEC
    index = PlanIndex(indexed, spec, snapshot_id="old")
    cross = {"rows": [{"row_index": 0, "matched_plan_key": "old:0"}]}
    attach_plan_facts(cross, index, [{"perf_comp": "X"}])
    assert len(cross["rows"][0]["plan_refs"]) == 3
    assert cross["rows"][0]["full_profile_quantity"] == 7
    assert cross["rows"][0]["plan_identity"]["profile_excel_o"] == "60,3"
    same_model = [{**items[0], "plan_key": str(i)} for i in range(501)]
    expanded = expand_entries(same_model, "old")
    ctx = plan_review.totals([plan_review.display_ref(r) for r in expanded["plan_refs"]])
    assert len(expanded["plan_refs"]) == 501
    assert ctx["nesta_folha"] == 2505
    assert not ctx["parcial"]
    copy = sheet()
    copy["cross_check"]["rows"][0].update(expanded)
    assert len(list(export_routes.facts_for(copy))) == 501


@pytest.mark.parametrize("quantity", [None, -1, "unknown"])
def test_unknown_or_negative_remaining_never_becomes_zero(quantity):
    result = expand_entries([{**entries()[0], "remaining_quantity": quantity}], "old")
    assert not result["plan_refs_valid"]
    assert result["plan_refs"][0]["assumed_quantity"] is None
    assert result["plan_line_meters"] is None


def test_unknown_length_keeps_quantity_without_fictitious_meters():
    result = expand_entries([{**entries()[0], "length_mm": None}], "old")
    assert result["full_profile_quantity"] == 5
    assert result["plan_line_meters"] is None


def test_legacy_recovery_uses_only_explicit_snapshot_and_never_replaces_zero(monkeypatch):
    source = sheet()
    source["cross_check"]["rows"] = [{"row_index": 0}]
    seen = []
    monkeypatch.setattr(loaders, "_fetch", lambda sql, params: [])
    monkeypatch.setattr(loaders, "plan_snapshot_info", lambda: pytest.fail("never use latest"))
    def order(sid, of):
        seen.append((sid, of))
        return entries()
    monkeypatch.setattr(plan_review, "fetch_order", order)
    prepared = export_source.prepare_sheets([source])[0]
    assert seen == [("old", "42")]
    assert [r["assumed_quantity"] for r in prepared["cross_check"]["rows"][0]["plan_refs"]] == [5, 2, 0]
    prepared["cross_check"]["rows"][0]["plan_refs"][0]["assumed_quantity"] = 0
    seen.clear()
    again = export_source.prepare_sheets([prepared])[0]
    assert again["cross_check"]["rows"][0]["plan_refs"][0]["assumed_quantity"] == 0
    assert not seen
    source["cross_check"].pop("snapshot_id")
    with pytest.raises(export_source.IncompleteExport) as error:
        export_source.prepare_sheets([source])
    assert error.value.problems[0]["row"] == 1


def test_legacy_children_are_preferred_to_plan_even_when_all_zero(monkeypatch):
    source = sheet()
    ref = source["cross_check"]["rows"][0].pop("plan_refs")[0]
    child = {**ref, "sheet_uid": source["uid"], "row_index": 0, "assumed_quantity": 0,
             "plan_snapshot_id": "old", "extra": {"plan_identity": ref}}
    monkeypatch.setattr(loaders, "_fetch", lambda *_: [child])
    monkeypatch.setattr(plan_review, "fetch_order", lambda *_: pytest.fail("archived zero wins"))
    result = export_source.prepare_sheets([source])[0]
    assert result["cross_check"]["rows"][0]["plan_refs"][0]["assumed_quantity"] == 0
    assert list(export_routes.facts_for(result)) == []


def test_archive_query_filters_source_and_reads_pg_only(monkeypatch):
    captured = {}
    def fetch(sql, params):
        captured.update(sql=sql, params=params)
        return [sheet(False)]
    monkeypatch.setattr(loaders, "_fetch", fetch)
    result = export_source.load_validated_sheets("2026-09-01", "2026-09-09", "ANA")
    assert result[0]["uid"] == "archive-only"
    assert "mes_kanban.validated_sheets" in captured["sql"]
    assert "source_app=%s" in captured["sql"] and "NOT LIKE" in captured["sql"]
    assert captured["params"] == (plan_review.SOURCE_APP, "2026-09-01", "2026-09-09", "ANA")


def test_export_failure_is_readable_and_never_local_fallback(client, monkeypatch):
    create_sheet(client)
    def unavailable(*args):
        raise ConnectionError("test")
    monkeypatch.setattr(export_source, "load_validated_sheets", unavailable)
    response = client.get("/export/basedados")
    assert response.status_code == 503
    assert response.headers["content-type"].startswith("text/html")
    problem = {"uid": "old-uid", "sheet_no": 99, "row": 3, "reason": "Snapshot desconhecido"}
    def incomplete(*args):
        raise export_source.IncompleteExport([problem])
    monkeypatch.setattr(export_source, "load_validated_sheets", incomplete)
    response = client.get("/export/basedados")
    assert response.status_code == 422
    assert "Snapshot desconhecido" in response.text and "99" in response.text


def test_validated_popup_opens_offline_without_index(client, monkeypatch):
    uid = insert_local(sheet())
    monkeypatch.setattr(loaders, "_fetch", lambda *_: pytest.fail("no PG"))
    monkeypatch.setattr(main, "get_index", lambda *_: pytest.fail("no index"))
    response = client.get(f"/sheet/{uid}/plano/0")
    assert response.status_code == 200
    assert "dados guardados na validação" in response.text
    assert "Feita nesta folha" in response.text
    assert "REF-2" in response.text


@pytest.mark.parametrize("engine", ["legacy", "v3"])
@pytest.mark.parametrize("full", [False, True])
def test_picker_changes_of_atomically_clears_blank_identity_and_preserves_production(client, monkeypatch, full, engine):
    from dataclasses import replace
    monkeypatch.setattr(main, "settings", replace(main.settings, cross_engine=engine))
    source = sheet(full)
    source["status"] = "review"
    source["sheet_data"]["rows"][0].update(duracao="00:42", n_corte_lote="LOTE-9")
    uid = insert_local(source)
    # Compare against the canonical facts actually stored by this template.
    conn = db.connect()
    try:
        original = deepcopy(db.get_sheet(conn, uid)["sheet_data"]["rows"][0])
    finally:
        conn.close()
    chosen = {**entries()[0], "production_order_no": "OF99", "sales_order_no": None, "customer_name": None,
              "component_ref": "NEW", "snapshot_id": "current", "plan_key": "new:key"}
    indexed = {**chosen, "of": "OF99", "ov": None, "cliente": None, "modelo": "NEW", "perfil": PROFILE,
               "comp_mm": chosen["length_mm"]}
    index = PlanIndex([indexed], loaders.PERFIS_SPEC if MTG2 else loaders.CANTONEIRAS_SPEC, snapshot_id="current")
    monkeypatch.setattr(loaders, "plan_snapshot_info", lambda: {"snapshot_id": "current"})
    monkeypatch.setattr(plan_review, "fetch_keys", lambda *_: [chosen])
    monkeypatch.setattr(main, "get_index", lambda loader: {} if loader == "load_employees" else index)
    payload = {"revision": get_revision(client, uid), "snapshot_id": "current", "selection_kind": "profile" if full else "reference",
               "plan_key": "new:key", "back": "/?status=pending&of=42&page=2"}
    response = client.post(f"/sheet/{uid}/rows/0/plan-selection", json=payload)
    assert response.status_code == 200, response.text
    conn = db.connect()
    try:
        saved = db.get_sheet(conn, uid)
        events = conn.execute("SELECT field_path,source,actor FROM edits WHERE sheet_uid=?", (uid,)).fetchall()
    finally:
        conn.close()
    row = saved["sheet_data"]["rows"][0]
    assert row["of"] == "99" and not row["ov"] and not row["cliente"]
    assert row["qtd"] == "4"
    for key in ("duracao", "n_corte_lote", "lote", "ultima_coluna"):
        assert row.get(key) == original.get(key)
    assert row.get("_plan_binding") is None if full else row["_plan_binding"]["plan_key"] == "new:key"
    assert row.get("perf_comp") == ("X" if full else None)
    assert any(e["source"] == "human" and e["actor"] == "plan-picker" for e in events)
    stale = client.post(f"/sheet/{uid}/rows/0/plan-selection", json=payload)
    assert stale.status_code == 409
    assert get_revision(client, uid) == saved["revision"]


def test_picker_rejects_old_snapshot_and_validated_sheet(client, monkeypatch):
    source = sheet(False)
    source["status"] = "review"
    uid = insert_local(source)
    monkeypatch.setattr(loaders, "plan_snapshot_info", lambda: {"snapshot_id": "new"})
    payload = {"revision": get_revision(client, uid), "snapshot_id": "old", "selection_kind": "reference", "plan_key": "old:0"}
    response = client.post(f"/sheet/{uid}/rows/0/plan-selection", json=payload)
    assert response.status_code == 409
    assert get_revision(client, uid) == payload["revision"]
    frozen_uid = insert_local(sheet(False))
    payload["revision"] = get_revision(client, frozen_uid)
    response = client.post(f"/sheet/{frozen_uid}/rows/0/plan-selection", json=payload)
    assert response.status_code == 409 and "só de leitura" in response.text


def test_validation_returns_to_history_context_and_sanitizes_destination(client):
    uid = create_sheet(client)
    edit(client, uid, "header.operador", "ANA")
    edit(client, uid, "header.data", "2026-09-01")
    history = "/?status=pending&operador=ANA&setor=Vanguard&data=2026-09-01&data_captura=2026-09-02&of=42&page=2"
    response = client.post(f"/sheet/{uid}/validate", data={"history_back": history, "back": "/captura"})
    assert response.status_code == 303
    url = urlsplit(response.headers["location"])
    assert url.path == "/" and "erro" not in parse_qs(url.query)
    params = parse_qs(url.query)
    for key, value in parse_qs(urlsplit(history).query).items():
        assert params[key] == value
    assert "validated" in params
    for hostile in ("//evil.test", "/\\evil.test", "/sheet/foo", "https://evil.test", "/%2fevil.test"):
        assert main._safe_history_back(hostile) is None


def test_pg_archive_wins_over_local_draft_and_basedados_never_needs_sqlite(client, monkeypatch):
    local = sheet(False)
    local["status"] = "review"
    local["sheet_data"]["rows"][0]["qtd"] = "99"
    uid = insert_local(local)
    archived = sheet(False)
    archived["uid"] = uid
    monkeypatch.setattr(export_source, "load_validated_sheets", lambda *_: [deepcopy(archived)])
    result = export_routes.export_sheets(main._conn, drafts=True)
    assert len(result) == 1
    assert list(export_routes.facts_for(result[0]))[0][1]["qtd"] == "4"
    assert len(export_routes.export_sheets(lambda: pytest.fail("BaseDados must not read SQLite"))) == 1


def test_historical_key_locates_snapshot_but_ambiguous_key_is_a_diagnostic(monkeypatch):
    legacy = sheet()
    legacy["cross_check"] = {"rows": [{"row_index": 0, "matched_plan_key": "old:0"}]}
    monkeypatch.setattr(loaders, "_fetch", lambda *_: [])
    monkeypatch.setattr(plan_review, "fetch_keys", lambda *_: [entries()[0]])
    called = []
    def order(snapshot, of):
        called.append((snapshot, of))
        return entries()
    monkeypatch.setattr(plan_review, "fetch_order", order)
    prepared = export_source.prepare_sheets([legacy])[0]
    assert called == [("old", "OF42")]
    assert len(prepared["cross_check"]["rows"][0]["plan_refs"]) == 3
    monkeypatch.setattr(plan_review, "fetch_keys", lambda *_: [entries()[0], {**entries()[0], "snapshot_id": "different"}])
    with pytest.raises(export_source.IncompleteExport):
        export_source.prepare_sheets([legacy])


def test_lookup_preserves_model_punctuation_source_scope_and_paginates_50(monkeypatch):
    calls = []
    def fetch(sql, params):
        calls.append((sql, params))
        if sql.startswith("SELECT 1"):
            return [{"exists": 1}] if "component_ref" in sql else []
        return [{"plan_key": str(i), "component_ref": "same-model"} for i in range(51)]
    monkeypatch.setattr(loaders, "_fetch", fetch)
    result = plan_review.lookup("old", "REF_%", include_done=True, offset=50)
    assert len(result["entries"]) == 50 and result["has_more"]
    assert result["offset"] == 50 and result["mode"] == "modelo"
    assert len(calls) == 4
    assert all(params[:2] == (plan_review.SOURCE_APP, "old") for _, params in calls)
    assert calls[-1][1][2] == "REF\\_\\%%"
    assert "LIMIT 51 OFFSET %s" in calls[-1][0]
    assert "remaining_quantity > 0" not in calls[-1][0]


@pytest.mark.parametrize("bad_key", [None, "", "old:0"])
def test_full_profile_rejects_missing_or_duplicate_keys(bad_key):
    result = expand_entries([entries()[0], {**entries()[1], "plan_key": bad_key}], "old")
    assert not result["plan_refs_valid"]
    assert result["full_profile_quantity"] is None
    assert result["plan_line_meters"] is None
    assert "chave" in result["plan_refs_error"]


def test_all_unknown_full_profile_meters_stay_unknown():
    item = {**entries()[0], "length_mm": None}
    indexed = {**item, "of": item["production_order_no"], "perfil": item["profile_type"],
               "modelo": item["component_ref"], "ov": item["sales_order_no"]}
    spec = loaders.PERFIS_SPEC if MTG2 else loaders.CANTONEIRAS_SPEC
    cross = {"rows": [{"row_index": 0, "matched_plan_key": item["plan_key"]}]}
    attach_plan_facts(cross, PlanIndex([indexed], spec, snapshot_id="old"), [{"perf_comp": "X"}])
    assert cross["summary"]["metros_teoricos"] is None
    assert cross["summary"]["metros_parciais"]
    assert cross["summary"]["desperdicio_m"] is None
