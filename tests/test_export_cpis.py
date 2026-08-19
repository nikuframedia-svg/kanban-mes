"""Export CPIS — o formato tabela da Metalogalva 2, verificado célula a célula."""

from __future__ import annotations

import datetime as dt
import io

from openpyxl import load_workbook

from app.web.export import CPIS_COLUMNS, build_cpis_workbook, cpis_filename_for, neutralize_xlsx
from tests.test_web import client, edit  # noqa: F401


def test_colunas_na_ordem_exata_do_original():
    labels = [label for _, label in CPIS_COLUMNS]
    assert labels == [
        "Data", "Cód. Funcionário", "Nome Funcionário", "Setor / Máquina Desc.",
        "Cód. Máquina", "OF", "OV", "Cliente", "Modelo", "QTD", "Qtd Metros",
        "M²", "Nesting", "Cesta Nº", "Duração", "Comprimento (mm)",
        "Largura (mm)", "Espessura (mm)", "CONI", "Nº Chapas",
        "Peso Consumido (t)", "Peso Produzido (t)", "Desperdício (t)",
        "% Desperdício", "Sucata", "Lote",
    ], "o contrato do ficheiro CPIS não pode desviar do original"


def test_workbook_formato_e_conteudo():
    rows = [{
        "data": dt.date(2026, 8, 10), "cod_funcionario": 47,
        "nome_funcionario": "JOSE SANTOS", "setor_maquina_desc": "Rapid 20T - 1",
        "of": "263322", "ov": "2504634", "cliente": "TENNET",
        "modelo": "AT2T562", "qtd": 4.0, "qtd_metros": 6.0,
        "comp_mm": 1500.0, "esp_mm": 5.0,
    }]
    wb = load_workbook(io.BytesIO(build_cpis_workbook(rows)))
    ws = wb.active
    assert ws.title == "Folha1"
    assert ws.freeze_panes == "A2"
    assert ws.cell(1, 1).value == "Data"
    assert ws.cell(1, 6).value == "OF"
    lida = ws.cell(2, 1).value          # o Excel guarda datas como datetime
    assert (lida.date() if isinstance(lida, dt.datetime) else lida) == dt.date(2026, 8, 10)
    assert ws.cell(2, 1).number_format == "DD-MM-YYYY"
    assert ws.cell(2, 6).value == "263322", "OF nua, sem prefixo"
    assert ws.cell(2, 10).value == 4.0
    assert ws.cell(2, 11).value == 6.0
    assert ws.cell(2, 18).number_format == "0.0", "espessura com 1 casa"
    assert ws.cell(1, 1).fill.start_color.rgb.endswith("2F5597")


def test_neutralize_anti_formula():
    assert neutralize_xlsx("=1+1") == "'=1+1"
    assert neutralize_xlsx("@cmd") == "'@cmd"
    assert neutralize_xlsx("TENNET") == "TENNET"
    assert neutralize_xlsx(4.0) == 4.0


def test_nome_do_ficheiro():
    assert cpis_filename_for(None, None, False) == "MigracaoNikufraCPIS_sempre.xlsx"
    assert cpis_filename_for("2026-08-10", "2026-08-10", False) == \
        "MigracaoNikufraCPIS_1-dia_2026-08-10.xlsx"
    assert cpis_filename_for("2026-08-01", "2026-08-10", True) == \
        "MigracaoNikufraCPIS_2026-08-01_2026-08-10_validadas.xlsx"


def test_rota_exporta_folha_real(client):  # noqa: F811
    """Fluxo completo: folha em staging → /export/cpis com filtros."""
    r = client.post("/upload", data={"template_name": "cantoneiras_kanban"})
    uid = r.headers["location"].rsplit("/", 1)[1]
    edit(client, uid, "header.operador", "Ze Manel")
    edit(client, uid, "header.data", "10/08/2026")
    edit(client, uid, "rows[0].of", "OF250001")
    edit(client, uid, "rows[0].qtd", "3")

    r = client.get("/export/cpis?de=2026-08-10&ate=2026-08-10")
    assert r.status_code == 200
    assert "MigracaoNikufraCPIS_1-dia_2026-08-10.xlsx" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 2, "uma linha de kanban = uma linha CPIS"
    assert ws.cell(2, 6).value == "250001", "OF do plano, nua"
    assert ws.cell(2, 3).value == "Ze Manel"
    # comprimento vem do plano (P1: comp_mm 1500) e os metros = qtd × comp
    assert ws.cell(2, 16).value == 1500.0
    assert ws.cell(2, 11).value == 4.5

    # fora do período → vazio
    r = client.get("/export/cpis?de=2026-01-01&ate=2026-01-31")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 1

    # só validadas → vazio (a folha está em revisão)
    r = client.get("/export/cpis?validadas=1")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 1
