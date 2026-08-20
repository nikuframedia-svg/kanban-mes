"""Export BaseDados — o formato Modelo_BaseDados_PerfisCantoneiras.xlsx,
comum aos dois setores kanban (11 colunas), verificado célula a célula."""

from __future__ import annotations

import datetime as dt
import io

from openpyxl import load_workbook

from app.web.export import (
    BASEDADOS_COLUMNS,
    basedados_filename_for,
    basedados_row_for,
    build_basedados_workbook,
)
from tests.test_web import client, edit  # noqa: F401


def test_colunas_na_ordem_exata_do_modelo():
    labels = [label for _, label in BASEDADOS_COLUMNS]
    assert labels == [
        "Data", "OperadorID", "Nome Operador", "Cod. Maquina", "Maquina",
        "OV", "OF", "Perfil", "Modelo", "Qtd [un.]", "Qtd [m]",
    ], "o contrato do ficheiro BaseDados não pode desviar do modelo"


def test_linha_com_valores_efetivos_e_refs_nuas():
    sheet = {"sheet_data": {"header": {
        "data": "10/08/2026", "operador": "Ze Manel", "n_operador": "47",
        "setor_maquina": "Rapid 20T - 1",
    }}}
    row = {"of": "OF263322", "ov": "", "perfil": "60x6", "modelo": "AT2T562",
           "qtd": "4"}
    cr = {"line_meters": 6.0,
          "cells": [{"field": "ov", "inherited": "OV2504634"}]}
    op = {"cod": 47, "name": "JOSE SANTOS"}
    bd = basedados_row_for(sheet, row, cr, op)
    assert bd["data"] == dt.date(2026, 8, 10)
    assert bd["operador_id"] == 47
    assert bd["nome_operador"] == "JOSE SANTOS"
    assert bd["cod_maquina"] is None, "sem tabela oficial de códigos, segue vazio"
    assert bd["maquina"] == "Rapid 20T - 1"
    assert bd["of"] == "263322", "OF nua, convenção do planeamento"
    assert bd["ov"] == "2504634", "herdada do bloco e nua"
    assert bd["perfil"] == "60x6"
    assert bd["modelo"] == "AT2T562"
    assert bd["qtd_un"] == 4.0
    assert bd["qtd_m"] == 6.0


def test_workbook_formato():
    rows = [{
        "data": dt.date(2026, 8, 10), "operador_id": 47,
        "nome_operador": "JOSE SANTOS", "cod_maquina": None,
        "maquina": "Rapid 20T - 1", "ov": "2504634", "of": "263322",
        "perfil": "60x6", "modelo": "=AT2T562", "qtd_un": 4.0, "qtd_m": 6.0,
    }]
    ws = load_workbook(io.BytesIO(build_basedados_workbook(rows))).active
    assert ws.title == "Folha1"
    assert ws.freeze_panes == "A2"
    assert [ws.cell(1, c).value for c in range(1, 12)] == [
        "Data", "OperadorID", "Nome Operador", "Cod. Maquina", "Maquina",
        "OV", "OF", "Perfil", "Modelo", "Qtd [un.]", "Qtd [m]",
    ]
    lida = ws.cell(2, 1).value
    assert (lida.date() if isinstance(lida, dt.datetime) else lida) == dt.date(2026, 8, 10)
    assert ws.cell(2, 1).number_format == "DD-MM-YYYY"
    assert ws.cell(2, 9).value == "'=AT2T562", "anti-fórmula neutralizada"
    assert ws.cell(2, 11).value == 6.0
    assert ws.cell(2, 11).number_format == "0.00"


def test_nome_do_ficheiro():
    assert basedados_filename_for(None, None, False) == "BaseDados_sempre.xlsx"
    assert basedados_filename_for("2026-08-10", "2026-08-10", False) == \
        "BaseDados_1-dia_2026-08-10.xlsx"
    assert basedados_filename_for("2026-08-01", "2026-08-10", True) == \
        "BaseDados_2026-08-01_2026-08-10_validadas.xlsx"


def test_rota_exporta_folha_real(client):  # noqa: F811
    """Fluxo completo: folha em staging → /export/basedados com filtros."""
    r = client.post("/upload", data={"template_name": "cantoneiras_kanban"})
    uid = r.headers["location"].rsplit("/", 1)[1]
    edit(client, uid, "header.operador", "Ze Manel")
    edit(client, uid, "header.data", "10/08/2026")
    edit(client, uid, "rows[0].of", "OF250001")
    edit(client, uid, "rows[0].qtd", "3")

    r = client.get("/export/basedados?de=2026-08-10&ate=2026-08-10")
    assert r.status_code == 200
    assert "BaseDados_1-dia_2026-08-10.xlsx" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 2, "uma linha de kanban = uma linha BaseDados"
    assert ws.cell(2, 7).value == "250001", "OF nua"
    assert ws.cell(2, 3).value == "Ze Manel"
    # metros = qtd × comprimento do plano (P1: comp_mm 1500 → 3 × 1.5 m)
    assert ws.cell(2, 11).value == 4.5

    # fora do período → vazio
    r = client.get("/export/basedados?de=2026-01-01&ate=2026-01-31")
    assert load_workbook(io.BytesIO(r.content)).active.max_row == 1
